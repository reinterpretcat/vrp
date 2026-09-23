#!/usr/bin/env python3
"""Build the roro.basic pragmatic-format VRP dataset from a RoRo order spreadsheet."""

from __future__ import annotations

import json
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path
from typing import NamedTuple

from openpyxl import load_workbook

from make_dataset import (
    MAPBOX_ATTEMPTS,
    MAPBOX_TIMEOUT,
    PROFILE,
    WASTE_FACILITY,
    build_routing_matrix,
    mapbox_token,
    read_json,
    unique_locations,
    write_json,
)

SERVICE_DURATION = 600.0
RELOAD_DURATION = 600.0
SHIFT_START_TIME = "09:00:00Z"
SHIFT_END_TIME = "18:00:00Z"
TRUCK_CAPACITY = 1
JOB_DEMAND = 1
TRUCK_IDS = [f"truck_{i}" for i in range(1, 5)]
TASK_TYPES = {
    "Emptying": "replacements",
    "Termination": "pickups",
    "Deployment": "deliveries",
}
MAPBOX_GEOCODE_URL = "https://api.mapbox.com/search/geocode/v6/forward"
MAPBOX_GEOCODE_PAUSE = 0.1

HERE = Path(__file__).resolve().parent
ORDERS_XLSX = HERE / "data" / "roro.basic.xlsx"
GEOCODES_PATH = HERE / "data" / "roro.basic.geocodes.json"
PROBLEM_PATH = HERE / "data" / "roro.basic.problem.json"
LOCATIONS_PATH = HERE / "data" / "roro.basic.locations.json"
MATRIX_PATH = HERE / "data" / "roro.basic.matrix.json"


class Order(NamedTuple):
    order_id: str
    task_type: str
    date: str
    address: str
    postcode: str
    place: str
    lat: float = 0.0
    lng: float = 0.0

    @property
    def query(self) -> str:
        return f"{self.address}, {self.postcode} {self.place}"


def load_orders(path: Path) -> list[Order]:
    sheet = load_workbook(path, read_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    orders: list[Order] = []
    for values in rows:
        row = dict(zip(header, values))
        if not row["Ordrenr"]:
            continue
        kind = row["Varenavn"].strip()
        if kind not in TASK_TYPES:
            raise SystemExit(f"Unknown order type {kind!r} in order {row['Ordrenr']}")
        date = row["Lev.dato"]
        orders.append(
            Order(
                order_id=str(row["Ordrenr"]),
                task_type=TASK_TYPES[kind],
                date=date.date().isoformat()
                if isinstance(date, datetime)
                else str(date),
                address=str(row["adr_nr"]).strip(),
                postcode=str(row["Anl.pnr"]).strip().zfill(4),
                place=str(row["Anl. sted"]).strip(),
            )
        )
    return orders


def request_geocode(params: dict) -> dict:
    url = f"{MAPBOX_GEOCODE_URL}?{urllib.parse.urlencode(params)}"
    last_error: Exception | None = None
    for attempt in range(MAPBOX_ATTEMPTS):
        try:
            with urllib.request.urlopen(url, timeout=MAPBOX_TIMEOUT) as response:
                return json.loads(response.read().decode())
        except urllib.error.HTTPError as err:
            body = err.read().decode(errors="replace")
            if err.code != 429:
                raise RuntimeError(
                    f"Mapbox Geocoding API HTTP {err.code}: {body}"
                ) from err
            last_error = RuntimeError(f"Mapbox Geocoding API rate limited: {body}")
        except urllib.error.URLError as err:
            last_error = RuntimeError(f"Mapbox Geocoding API request failed: {err}")
        time.sleep(2**attempt)
    raise last_error or RuntimeError("Mapbox Geocoding API request failed")


def geocode(order: Order, token: str) -> dict:
    payload = request_geocode(
        {
            "q": order.query,
            "country": "no",
            "types": "address",
            "limit": 1,
            "access_token": token,
        }
    )
    features = payload.get("features") or []
    if not features:
        raise RuntimeError(f"Mapbox found no match for {order.query!r}")
    feature = features[0]
    lng, lat = feature["geometry"]["coordinates"]
    properties = feature.get("properties", {})
    return {
        "lat": lat,
        "lng": lng,
        "match": properties.get("full_address") or properties.get("name"),
        "confidence": properties.get("match_code", {}).get("confidence"),
    }


def geocode_orders(orders: list[Order]) -> list[Order]:
    cache = read_json(GEOCODES_PATH)
    geocodes: dict[str, dict] = cache if isinstance(cache, dict) else {}
    token: str | None = None
    for order in orders:
        if order.query in geocodes:
            continue
        token = token or mapbox_token()
        result = geocode(order, token)
        print(
            f"Geocoded {order.query!r} → {result['match']!r} ({result['confidence']})"
        )
        geocodes[order.query] = result
        time.sleep(MAPBOX_GEOCODE_PAUSE)
    write_json(GEOCODES_PATH, geocodes)
    return [
        order._replace(
            lat=geocodes[order.query]["lat"], lng=geocodes[order.query]["lng"]
        )
        for order in orders
    ]


def job(order: Order) -> dict:
    return {
        "id": order.order_id,
        order.task_type: [
            {
                "places": [
                    {
                        "location": {"lat": order.lat, "lng": order.lng},
                        "duration": SERVICE_DURATION,
                    }
                ],
                "demand": [JOB_DEMAND],
            }
        ],
    }


def fleet(date: str, reload_count: int) -> dict:
    return {
        "vehicles": [
            {
                "typeId": "truck",
                "vehicleIds": TRUCK_IDS,
                "profile": {"matrix": PROFILE},
                "costs": {
                    "fixed": 22.0,
                    "distance": 0.0002,
                    "time": 0.004806,
                },
                "shifts": [
                    {
                        "start": {
                            "earliest": f"{date}T{SHIFT_START_TIME}",
                            "location": dict(WASTE_FACILITY),
                        },
                        "end": {
                            "latest": f"{date}T{SHIFT_END_TIME}",
                            "location": dict(WASTE_FACILITY),
                        },
                        # each reload can be visited at most once, so allow one per job
                        "reloads": [
                            {
                                "location": dict(WASTE_FACILITY),
                                "duration": RELOAD_DURATION,
                            }
                            for _ in range(reload_count)
                        ],
                    }
                ],
                "capacity": [TRUCK_CAPACITY],
            }
        ],
        "profiles": [{"name": PROFILE}],
    }


def build_problem(orders: list[Order]) -> dict:
    dates = {order.date for order in orders}
    if len(dates) != 1:
        raise SystemExit(f"Expected orders for a single date, got {sorted(dates)}")
    jobs = [job(order) for order in orders]
    return {"plan": {"jobs": jobs}, "fleet": fleet(dates.pop(), len(jobs))}


def cached_matrix(locations: list[dict]) -> dict | None:
    if read_json(LOCATIONS_PATH) != json.loads(json.dumps(locations)):
        return None
    matrix = read_json(MATRIX_PATH)
    expected = len(locations) ** 2
    if not isinstance(matrix, dict) or matrix.get("profile") != PROFILE:
        return None
    if len(matrix.get("travelTimes") or []) != expected:
        return None
    if len(matrix.get("distances") or []) != expected:
        return None
    return matrix


def routing_matrix(locations: list[dict]) -> dict:
    matrix = cached_matrix(locations)
    if matrix is None:
        return build_routing_matrix(locations, mapbox_token())
    print(f"Reusing cached {len(locations)}×{len(locations)} matrix from {MATRIX_PATH}")
    return matrix


def main() -> None:
    orders = geocode_orders(load_orders(ORDERS_XLSX))
    locations = unique_locations(orders)
    matrix = routing_matrix(locations)
    write_json(PROBLEM_PATH, build_problem(orders))
    write_json(LOCATIONS_PATH, locations)
    write_json(MATRIX_PATH, matrix)
    print(f"Wrote {len(orders)} jobs and {len(TRUCK_IDS)} trucks to {PROBLEM_PATH}")
    print(f"Wrote {len(locations)} unique locations to {LOCATIONS_PATH}")
    print(f"Wrote {len(locations)}×{len(locations)} routing matrix to {MATRIX_PATH}")


if __name__ == "__main__":
    main()
